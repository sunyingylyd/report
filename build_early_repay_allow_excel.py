#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""允许/不允许提前还款：到期周度+月度，金额 USD，写本地 Excel。"""
from __future__ import annotations

import os
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SQL_USD = HERE / "appendix_early_repay_allow_usd.sql"
SQL_MXN = HERE / "appendix_early_repay_allow_mxn.sql"
SQL_MXN_RECORD = HERE / "appendix_early_repay_allow_mxn_record.sql"
SQL_MXN_RECORD_EXPECT = HERE / "appendix_early_repay_allow_mxn_record_expect.sql"
XLSX_USD = HERE / "提前还款允许组对照_到期20260101-20260920.xlsx"
XLSX_MXN = HERE / "提前还款允许组对照_到期20260101-20260920_比索.xlsx"
XLSX_MXN_RECORD = HERE / "提前还款允许组对照_到期20260101-20260920_比索_还款流水.xlsx"
XLSX_MXN_RECORD_EXPECT = HERE / "提前还款允许组对照_到期20260101-20260920_比索_还款流水_入催应还.xlsx"
FLOW_TABLE = "repayment_plan_record_copy"
COLL_AMT_FIELD = "apply_amt_period"
SQL_PATH = SQL_USD
XLSX_PATH = XLSX_USD
CURRENCY = "USD"
AMT_KIND = "usd"

METRICS = [
    ("到期订单数", "n_due", "int"),
    ("提前结清订单数", "n_early", "int"),
    ("订单提前结清率", "early_n_rate", "pct"),
    ("入催订单数", "n_coll", "int"),
    ("订单入催率", "coll_n_rate", "pct"),
    ("实时逾期订单数", "n_st8", "int"),
    ("订单实时逾期率", "st8_n_rate", "pct"),
    ("到期订单金额", "due_amt", "usd"),
    ("提前结清订单金额", "early_amt", "usd"),
    ("金额提前结清率", "early_amt_rate", "pct"),
    ("入催订单金额", "coll_amt", "usd"),
    ("金额入催率", "coll_amt_rate", "pct"),
    ("实时逾期金额", "st8_amt", "usd"),
    ("金额实时逾期率", "st8_amt_rate", "pct"),
    ("盈利率", "profit_flow", "pct"),
    ("盈利率1", "profit_plan", "pct"),
]


def connect():
    if not os.environ.get("PGPASSWORD"):
        raise SystemExit("请先 source ~/.t7_dashboard.env")
    import psycopg2

    conn = psycopg2.connect(
        host=os.environ.get("PGHOST", "47.89.225.85"),
        port=int(os.environ.get("PGPORT", "8000")),
        dbname=os.environ.get("PGDATABASE", "kaby_dw"),
        user=os.environ.get("PGUSER", "wangchuanliang_readonly"),
        password=os.environ["PGPASSWORD"],
        connect_timeout=30,
        options="-c statement_timeout=0",
    )
    conn.autocommit = True
    return conn


def exec_sql(conn, sql: str):
    cur = conn.cursor()
    cur.execute(sql)
    cur.close()


def run_ddl(conn):
    sql_text = SQL_PATH.read_text(encoding="utf-8")
    for stmt in sql_text.split(";"):
        s = stmt.strip()
        if not s:
            continue
        body = "\n".join(ln for ln in s.splitlines() if not ln.strip().startswith("--")).strip()
        if not body:
            continue
        print("SQL:", body.split()[0], body.split()[1] if len(body.split()) > 1 else "", flush=True)
        exec_sql(conn, s)
        print("  done", flush=True)


def fetch_agg(conn) -> pd.DataFrame:
    cur = conn.cursor()
    cur.execute(
        """
        SELECT scope, grain, period_d, allow_flag,
               n_due, n_early, n_coll, n_st8,
               due_amt, early_amt, coll_amt, st8_amt, flow_usd, repaid_usd
        FROM tmp_agg
        """
    )
    cols = [d[0] for d in cur.description]
    rows = cur.fetchall()
    cur.close()
    df = pd.DataFrame(rows, columns=cols)
    df["period_d"] = pd.to_datetime(df["period_d"]).dt.date
    for c in [
        "n_due", "n_early", "n_coll", "n_st8",
        "due_amt", "early_amt", "coll_amt", "st8_amt", "flow_usd", "repaid_usd",
    ]:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    return df


def rate(num, den):
    den = float(den or 0)
    if den == 0:
        return None
    return float(num) / den


def metrics_from_row(r) -> dict:
    n_due = int(r["n_due"])
    n_early = int(r["n_early"])
    n_coll = int(r["n_coll"])
    n_st8 = int(r["n_st8"])
    due_amt = float(r["due_amt"])
    early_amt = float(r["early_amt"])
    coll_amt = float(r["coll_amt"])
    st8_amt = float(r["st8_amt"])
    flow = float(r["flow_usd"])
    repaid = float(r["repaid_usd"])
    return {
        "n_due": n_due,
        "n_early": n_early,
        "early_n_rate": rate(n_early, n_due),
        "n_coll": n_coll,
        "coll_n_rate": rate(n_coll, n_due),
        "n_st8": n_st8,
        "st8_n_rate": rate(n_st8, n_due),
        "due_amt": due_amt,
        "early_amt": early_amt,
        "early_amt_rate": rate(early_amt, due_amt),
        "coll_amt": coll_amt,
        "coll_amt_rate": rate(coll_amt, due_amt),
        "st8_amt": st8_amt,
        "st8_amt_rate": rate(st8_amt, due_amt),
        "profit_flow": (flow / due_amt - 1.0) if due_amt else None,
        "profit_plan": (repaid / due_amt - 1.0) if due_amt else None,
    }


def empty_metrics() -> dict:
    return {k: None if t != "int" else 0 for _, k, t in METRICS}


def period_label(grain: str, d: date) -> str:
    if grain == "month":
        return f"{d.year}年{d.month}月"
    sun = d + timedelta(days=6)
    return f"{d.strftime('%Y-%m-%d')}~{sun.strftime('%m-%d')}"


def write_pivot(ws, agg: pd.DataFrame, grain: str, title: str):
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    fill_h = PatternFill("solid", fgColor="1F4E79")
    fill_deny = PatternFill("solid", fgColor="F4F4F4")
    fill_allow = PatternFill("solid", fgColor="E2EFDA")
    font_h = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    font_h2 = Font(bold=True, name="Calibri", size=10)
    font_m = Font(name="Calibri", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sub = agg[agg["grain"] == grain].copy()
    periods = sorted(sub["period_d"].unique())

    ws["A1"] = title
    ws["A1"].font = Font(bold=True, name="Calibri", size=13)
    last_col = max(1 + 2 * len(periods), 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)

    ws["A2"] = "指标"
    ws["A2"].fill = fill_h
    ws["A2"].font = font_h
    ws["A2"].alignment = center
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)

    by = {}
    for _, row in sub.iterrows():
        by[(row["period_d"], row["allow_flag"])] = metrics_from_row(row)

    for i, pval in enumerate(periods):
        c0 = 2 + i * 2
        ws.cell(2, c0, period_label(grain, pval)).fill = fill_h
        ws.cell(2, c0).font = font_h
        ws.cell(2, c0).alignment = center
        ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + 1)
        ws.cell(3, c0, "不允许提前还款用户").fill = fill_deny
        ws.cell(3, c0 + 1, "允许提前还款用户").fill = fill_allow
        ws.cell(3, c0).font = font_h2
        ws.cell(3, c0 + 1).font = font_h2
        ws.cell(3, c0).alignment = center
        ws.cell(3, c0 + 1).alignment = center

    for r, (label, field, kind) in enumerate(METRICS, start=4):
        cell = ws.cell(r, 1, label)
        cell.font = font_m
        cell.border = thin
        for i, pval in enumerate(periods):
            for j, flag in enumerate(("deny", "allow")):
                m = by.get((pval, flag)) or empty_metrics()
                val = m.get(field)
                c = ws.cell(r, 2 + i * 2 + j, val)
                c.font = font_m
                c.border = thin
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.fill = fill_deny if flag == "deny" else fill_allow
                if val is None:
                    pass
                elif kind == "int":
                    c.number_format = "#,##0"
                elif kind == "usd":
                    c.number_format = "#,##0.00" if CURRENCY == "USD" else "#,##0"
                elif kind == "pct":
                    c.number_format = "0.00%"

    ws.column_dimensions["A"].width = 20
    for col in range(2, 2 + max(2 * len(periods), 1)):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 28
    ws.freeze_panes = "B4"


def write_口径(ws):
    if CURRENCY == "USD":
        lines = [
            "口径",
            "样本：order_loan_f_v2_copy，is_remit=1 且 remit_amt>0，due_date 在 2026-01-01～2026-09-20",
            "月份：due_date 所在自然月；周度：due_date 所在周一～周日（DATE_TRUNC week）",
            "firstloan / reloan 为订单标签。允许提前还款：首贷用户ID末位4-9；续贷末位1-9；其余不允许",
            "提前结清：repaid_date>2000-01-01 且 repaid_date<due_date",
            "入催件数：overdue_days>0；实时逾期件数：loan_status_code=8",
            "到期/提前结清金额：订单 apply_amt × remit_date 汇率（一单一份，不用分期 apply_amt）",
            "入催金额：overdue_days_period>0 的 SUM(apply_amt_period)×放款日汇率",
            "实时逾期金额：loan_status_code_period=8 的 SUM(apply_amt_period)×放款日汇率",
            "汇率：exchange_rate_copy country=MXN，local2usd，snap_date=业务日；缺日用不超过该日的最近一条",
            "盈利率（主）：repayment_plan_record_copy.amount（已扣折扣）× payin_date 汇率 / 到期金额USD − 1",
            "盈利率1：repaid_amt_period（未扣折扣）×（repaid_date_period>2000-01-01 用该日，否则 due_date_period）汇率 / 到期金额USD − 1",
            "金额单位：美元",
        ]
    else:
        lines = [
            "口径",
            "样本：order_loan_f_v2_copy，is_remit=1 且 remit_amt>0，due_date 在 2026-01-01～2026-09-20",
            "月份：due_date 所在自然月；周度：due_date 所在周一～周日（DATE_TRUNC week）",
            "firstloan / reloan 为订单标签。允许提前还款：首贷用户ID末位4-9；续贷末位1-9；其余不允许",
            "提前结清：repaid_date>2000-01-01 且 repaid_date<due_date",
            "入催件数：overdue_days>0；实时逾期件数：loan_status_code=8",
            "到期/提前结清金额：订单 apply_amt（一单一份，不用分期 apply_amt）",
            "入催金额：overdue_days_period>0 的 SUM(" + COLL_AMT_FIELD + ")",
            "实时逾期金额：loan_status_code_period=8 的 SUM(apply_amt_period)",
            "不换汇，金额为墨西哥比索原币",
            f"盈利率（主）：{FLOW_TABLE}.amount（已扣折扣）/ 到期金额 − 1",
            "盈利率1：SUM(repaid_amt_period)（未扣折扣）/ 到期金额 − 1",
            "金额单位：墨西哥比索 MXN",
        ]
    ws["A1"] = "口径说明"
    ws["A1"].font = Font(bold=True, size=13)
    for i, t in enumerate(lines, start=2):
        ws.cell(i, 1, t)
        ws.column_dimensions["A"].width = 120


def main():
    global SQL_PATH, XLSX_PATH, CURRENCY, FLOW_TABLE, COLL_AMT_FIELD
    import sys

    if "--mxn-record-expect" in sys.argv:
        SQL_PATH = SQL_MXN_RECORD_EXPECT
        XLSX_PATH = XLSX_MXN_RECORD_EXPECT
        CURRENCY = "MXN"
        FLOW_TABLE = "repayment_record_copy"
        COLL_AMT_FIELD = "expect_repay_amt_period"
    elif "--mxn-record" in sys.argv:
        SQL_PATH = SQL_MXN_RECORD
        XLSX_PATH = XLSX_MXN_RECORD
        CURRENCY = "MXN"
        FLOW_TABLE = "repayment_record_copy"
        COLL_AMT_FIELD = "apply_amt_period"
    elif "--mxn" in sys.argv:
        SQL_PATH = SQL_MXN
        XLSX_PATH = XLSX_MXN
        CURRENCY = "MXN"
        FLOW_TABLE = "repayment_plan_record_copy"
        COLL_AMT_FIELD = "apply_amt_period"
    else:
        SQL_PATH = SQL_USD
        XLSX_PATH = XLSX_USD
        CURRENCY = "USD"
        FLOW_TABLE = "repayment_plan_record_copy"
        COLL_AMT_FIELD = "apply_amt_period"
    conn = connect()
    print("loading…", CURRENCY, FLOW_TABLE, flush=True)
    run_ddl(conn)
    df = fetch_agg(conn)
    conn.close()
    print("agg rows", len(df), flush=True)
    wb = Workbook()
    sheets = [
        ("全部_月度", "all", "month"),
        ("首贷_月度", "firstloan", "month"),
        ("续贷_月度", "reloan", "month"),
        ("全部_周度", "all", "week"),
        ("首贷_周度", "firstloan", "week"),
        ("续贷_周度", "reloan", "week"),
    ]
    first = True
    unit = "金额USD" if CURRENCY == "USD" else "金额墨西哥比索"
    for name, scope, grain in sheets:
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        sub = df[df["scope"] == scope]
        title = {
            "all": "全部订单",
            "firstloan": "首贷订单 firstloan",
            "reloan": "续贷订单 reloan",
        }[scope] + (" · 月度" if grain == "month" else " · 周度（周一～周日）")
        title += f" · 到期 2026-01-01～2026-09-20 · {unit}"
        write_pivot(ws, sub, grain, title)
    ws = wb.create_sheet("口径")
    write_口径(ws)
    wb.save(XLSX_PATH)
    print("saved", XLSX_PATH, flush=True)


if __name__ == "__main__":
    main()
